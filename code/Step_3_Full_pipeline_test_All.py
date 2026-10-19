# -*- coding: utf-8 -*-

import os
import re
import numpy as np
import nibabel as nib
import cv2
import tensorflow as tf
import tensorflow.keras.backend as K
import joblib
import pandas as pd
from glob import glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── PARAMETRELER ─────────────────────────────────────────
SEG_MODEL_PATH = r"E:\3D_sinus\models\best_sinus_model.h5"
MODEL_YAS_PATH = r"E:\3D_sinus\ml_data\model_yas.joblib"
MODEL_CIN_PATH = r"E:\3D_sinus\ml_data\model_cinsiyet.joblib"
EXCEL_REF_PATH = r"E:\3D_sinus\Uzman_kontrol_maske_hacmi2.xlsx"
TEST_IMG_DIR   = r"E:\3D_sinus\test_nifti\img"
TEST_MASK_DIR  = r"E:\3D_sinus\test_nifti\mask"
OUTPUT_EXCEL   = r"E:\3D_sinus\results\toplu_test_sonuclari.xlsx"

OFFSET_H, OFFSET_W = 83, 93
TARGET_H, TARGET_W = 261, 324
MODEL_H, MODEL_W   = 256, 256
W_CENTER, W_WIDTH  = 400, 2000
FEATURE_COLS       = ['maxilla_R', 'maxilla_L', 'maxilla_T']
# ──────────────────────────────────────────────────────────


def f1_score(y_true, y_pred):
    y_true_f = K.flatten(y_true)
    y_pred_f = K.flatten(y_pred)
    inter    = K.sum(y_true_f * y_pred_f)
    return (2. * inter + K.epsilon()) / (K.sum(y_true_f) + K.sum(y_pred_f) + K.epsilon())


def apply_windowing(image, wc, ww):
    lo = wc - ww / 2.0
    hi = wc + ww / 2.0
    return (((np.clip(image, lo, hi) - lo) / (hi - lo)) * 255.0).astype(np.uint8)


def calculate_metrics(y_true, y_pred):
    yt    = y_true.flatten().astype(np.float32)
    yp    = y_pred.flatten().astype(np.float32)
    inter = np.sum(yt * yp)
    dice  = (2. * inter) / (np.sum(yt) + np.sum(yp) + 1e-7)
    iou   = inter / (np.sum(yt) + np.sum(yp) - inter + 1e-7)
    prec  = inter / (np.sum(yp) + 1e-7)
    rec   = inter / (np.sum(yt) + 1e-7)
    return dice, iou, prec, rec


def parse_filename(filename):
    """
    '121-7.nii.gz'  -> key='121-7'
    'k10-3.nii.gz'  -> key='k10-3'
    """
    base = os.path.basename(filename)
    for ext in ['.nii.gz', '.nii']:
        base = base.replace(ext, '')
    return base


def load_excel_ref(excel_path):
    """
    Uzman_kontrol_maske_hacmi2.xlsx dosyasini okur.
    Doner: { 'hasta_adi_key': { uzman_R, uzman_L, uzman_T, yas, cinsiyet } }
    """
    df = pd.read_excel(excel_path)
    # Sutun isimlerini normalize et
    df.columns = [c.strip() for c in df.columns]
    df = df.rename(columns={
        'Hasta Adı'    : 'hasta_adi',
        'maxilla R ıtk': 'uzman_R',
        'maxilla L ıtk': 'uzman_L',
        'maxilla T'    : 'uzman_T',
        'yas'          : 'yas',
        'cinsiyet'     : 'cinsiyet',
    })
    df['hasta_adi'] = df['hasta_adi'].astype(str).str.strip()
    lookup = {}
    for _, row in df.iterrows():
        cin = str(row.get('cinsiyet', '')).strip().lower()
        if cin in ('nan', ''):
            cin = 'bilinmiyor'
        lookup[row['hasta_adi']] = {
            'uzman_R' : float(row['uzman_R']) if pd.notna(row['uzman_R']) else None,
            'uzman_L' : float(row['uzman_L']) if pd.notna(row['uzman_L']) else None,
            'uzman_T' : float(row['uzman_T']) if pd.notna(row['uzman_T']) else None,
            'yas'     : int(row['yas'])        if pd.notna(row['yas'])    else None,
            'cinsiyet': cin,
        }
    print(f"[BILGI] Referans Excel'den {len(lookup)} hasta yuklendi.")
    return lookup


def segment_sag_sol(seg_model, img_path, mask_path=None):
    """
    Segmentasyon yapar.
    Maskeyi orta eksenden (sutun 256) bolerek sag/sol hacmi ayri hesaplar.
    """
    img_can  = nib.as_closest_canonical(nib.load(img_path))
    img_data = img_can.get_fdata()
    vx, vy, vz   = img_can.header.get_zooms()[:3]
    unit_vol_mm3 = vx * vy * vz

    mask_data = None
    if mask_path and os.path.exists(mask_path):
        mask_data = nib.as_closest_canonical(nib.load(mask_path)).get_fdata()

    pred_list, true_list = [], []

    for i in range(img_data.shape[2]):
        if mask_data is not None:
            raw_mask = (mask_data[:, :, i] > 0).astype(np.uint8)
            if not np.any(raw_mask):
                continue
        else:
            raw_mask = None

        aligned  = np.rot90(img_data[:, :, i], k=1)
        windowed = apply_windowing(aligned, W_CENTER, W_WIDTH)
        cropped  = windowed[OFFSET_H:OFFSET_H+TARGET_H, OFFSET_W:OFFSET_W+TARGET_W]
        inp      = cv2.resize(cropped, (MODEL_W, MODEL_H), interpolation=cv2.INTER_AREA)
        inp      = np.expand_dims(inp.astype(np.float32) / 255.0, axis=(0, -1))

        pred256  = (seg_model.predict(inp, verbose=0)[0, :, :, 0] > 0.5).astype(np.uint8)

        if raw_mask is None and not np.any(pred256):
            continue

        canvas = np.zeros((512, 512), dtype=np.uint8)
        canvas[OFFSET_H:OFFSET_H+TARGET_H, OFFSET_W:OFFSET_W+TARGET_W] = \
            cv2.resize(pred256, (TARGET_W, TARGET_H), interpolation=cv2.INTER_NEAREST)

        pred_list.append(canvas)
        if raw_mask is not None:
            true_list.append(np.rot90(raw_mask, k=1))

    if not pred_list:
        return None

    pred_vol   = np.stack(pred_list, axis=-1)   # (512, 512, n_slice)
    mid        = pred_vol.shape[1] // 2
    pred_R_cm3 = np.sum(pred_vol[:, mid:, :]) * unit_vol_mm3 / 1000.0
    pred_L_cm3 = np.sum(pred_vol[:, :mid, :]) * unit_vol_mm3 / 1000.0
    pred_T_cm3 = np.sum(pred_vol)              * unit_vol_mm3 / 1000.0

    result = {
        'pred_R_cm3'   : round(pred_R_cm3, 4),
        'pred_L_cm3'   : round(pred_L_cm3, 4),
        'pred_T_cm3'   : round(pred_T_cm3, 4),
        'n_aktif_kesit': len(pred_list),
        'slice_thick'  : round(float(vz), 4),
    }

    if true_list:
        true_vol   = np.stack(true_list, axis=-1)
        true_T_cm3 = np.sum(true_vol) * unit_vol_mm3 / 1000.0
        dice, iou, prec, rec = calculate_metrics(true_vol, pred_vol)
        result.update({
            'gt_T_cm3' : round(true_T_cm3, 4),
            'dice'     : round(dice, 4),
            'iou'      : round(iou,  4),
            'precision': round(prec, 4),
            'recall'   : round(rec,  4),
        })

    return result


def predict_age_gender(seg_result, yas_model, cin_model):
    X          = np.array([[seg_result['pred_R_cm3'],
                             seg_result['pred_L_cm3'],
                             seg_result['pred_T_cm3']]])
    yas_tahmin = float(yas_model.predict(X)[0])
    cin_enc    = int(cin_model.predict(X)[0])
    cin_prob   = cin_model.predict_proba(X)[0]
    return {
        'yas_tahmin'   : round(yas_tahmin, 1),
        'cinsiyet'     : 'erkek' if cin_enc == 1 else 'kadin',
        'cin_guven_pct': round(max(cin_prob) * 100, 1),
    }


def pct_error(pred, gercek):
    if pred is not None and gercek and gercek != 0:
        return round(abs(pred - gercek) / gercek * 100, 2)
    return None


# ── Excel yazıcı ─────────────────────────────────────────

def write_excel(rows, output_path):

    # Stiller
    THIN       = Side(style='thin', color='CCCCCC')
    BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    HDR_FONT   = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    DATA_FONT  = Font(name='Arial', size=9)
    HDR_FILL   = PatternFill('solid', fgColor='1F4E79')
    FILL_BLUE  = PatternFill('solid', fgColor='D6E4F0')   # yas grubu
    FILL_GREEN = PatternFill('solid', fgColor='D8F0DC')   # cinsiyet grubu
    FILL_ORG   = PatternFill('solid', fgColor='FFF3E0')   # hacim grubu
    FILL_PUR   = PatternFill('solid', fgColor='F3E5F5')   # seg metrikleri
    FILL_GRAY  = PatternFill('solid', fgColor='F5F5F5')   # diger
    FILL_WHITE = PatternFill('solid', fgColor='FFFFFF')
    FILL_ALT   = PatternFill('solid', fgColor='F9FAFB')

    # Sütun tanımı: (başlık, veri_key, format, fill)
    cols = [
        ('Hasta Adi',          'hasta_adi',     'text',   FILL_GRAY),
        ('Gercek Yas',         'yas',            'int',    FILL_BLUE),
        ('Tahmin Yas',         'yas_tahmin',     'f1',     FILL_BLUE),
        ('Yas Hatasi (yil)',   'yas_hata',       'f1',     FILL_BLUE),
        ('Yas Hata %',         'yas_hata_pct',   'pct',    FILL_BLUE),
        ('Gercek Cinsiyet',    'cinsiyet',       'text',   FILL_GREEN),
        ('Tahmin Cinsiyet',    'tahmin_cin',     'text',   FILL_GREEN),
        ('Cinsiyet Dogru mu',  'cin_dogru',      'text',   FILL_GREEN),
        ('Cin. Guven %',       'cin_guven_pct',  'pct',    FILL_GREEN),
        ('Uzman R (cm3)',      'uzman_R',        'f4',     FILL_ORG),
        ('Pred R (cm3)',       'pred_R_cm3',     'f4',     FILL_ORG),
        ('R Hata %',           'R_hata_pct',     'pct',    FILL_ORG),
        ('Uzman L (cm3)',      'uzman_L',        'f4',     FILL_ORG),
        ('Pred L (cm3)',       'pred_L_cm3',     'f4',     FILL_ORG),
        ('L Hata %',           'L_hata_pct',     'pct',    FILL_ORG),
        ('Uzman T (cm3)',      'uzman_T',        'f4',     FILL_ORG),
        ('Pred T (cm3)',       'pred_T_cm3',     'f4',     FILL_ORG),
        ('T Hata %',           'T_hata_pct',     'pct',    FILL_ORG),
        ('Dice',               'dice',           'f4',     FILL_PUR),
        ('IoU',                'iou',            'f4',     FILL_PUR),
        ('Precision',          'precision',      'f4',     FILL_PUR),
        ('Recall',             'recall',         'f4',     FILL_PUR),
        ('Aktif Kesit',        'n_aktif_kesit',  'int',    FILL_GRAY),
        ('Kesit Kalinligi mm', 'slice_thick',    'f4',     FILL_GRAY),
    ]

    fmt_map = {'f4': '0.0000', 'f1': '0.0', 'pct': '0.00"%"',
               'int': '0', 'text': '@'}

    wb = Workbook()

    # ════════════════════════════════
    # SAYFA 1: TAHMİNLER
    # ════════════════════════════════
    ws = wb.active
    ws.title = 'Tahminler'

    # Başlık satırı
    for ci, (header, _, _, fill) in enumerate(cols, 1):
        c = ws.cell(row=1, column=ci, value=header)
        c.font      = HDR_FONT
        c.fill      = HDR_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border    = BORDER
    ws.row_dimensions[1].height = 32

    # Hataları hesapla ve satırları yaz
    for ri, row in enumerate(rows, 2):
        gy = row.get('yas')
        ty = row.get('yas_tahmin')
        gc = str(row.get('cinsiyet', '')).lower().strip()
        tc = str(row.get('tahmin_cin', '')).lower().strip()

        row['yas_hata']     = round(abs(ty - gy), 1)       if (ty is not None and gy is not None) else None
        row['yas_hata_pct'] = pct_error(ty, gy)
        row['R_hata_pct']   = pct_error(row.get('pred_R_cm3'), row.get('uzman_R'))
        row['L_hata_pct']   = pct_error(row.get('pred_L_cm3'), row.get('uzman_L'))
        row['T_hata_pct']   = pct_error(row.get('pred_T_cm3'), row.get('uzman_T'))
        row['cin_dogru']    = 'EVET' if gc == tc else ('?' if gc == 'bilinmiyor' else 'HAYIR')

        zebra = FILL_WHITE if ri % 2 == 0 else FILL_ALT

        for ci, (header, key, fmt, grp_fill) in enumerate(cols, 1):
            val = row.get(key)
            c   = ws.cell(row=ri, column=ci, value=val)
            c.font      = DATA_FONT
            c.border    = BORDER
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.fill      = grp_fill if grp_fill else zebra
            if val is not None and fmt in fmt_map:
                c.number_format = fmt_map[fmt]

            # Cinsiyet doğru/yanlış rengi
            if key == 'cin_dogru':
                if val == 'EVET':
                    c.font = Font(name='Arial', size=9, bold=True, color='1B5E20')
                elif val == 'HAYIR':
                    c.font = Font(name='Arial', size=9, bold=True, color='B71C1C')

    # Sütun genişlikleri
    widths = [12,11,11,16,11,15,15,16,12, 13,13,10, 13,13,10, 13,13,10, 10,10,10,10, 12,16]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"

    # ════════════════════════════════
    # SAYFA 2: ÖZET
    # ════════════════════════════════
    ws2 = wb.create_sheet('Ozet')
    ws2.column_dimensions['A'].width = 30
    ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 16
    ws2.column_dimensions['D'].width = 16

    def sec_hdr(row, title):
        ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        c = ws2.cell(row=row, column=1, value=title)
        c.font      = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        c.fill      = PatternFill('solid', fgColor='1F4E79')
        c.alignment = Alignment(horizontal='center')
        return row + 1

    def col_hdr(row, labels):
        for ci, lbl in enumerate(labels, 1):
            c = ws2.cell(row=row, column=ci, value=lbl)
            c.font      = Font(name='Arial', bold=True, size=10)
            c.fill      = PatternFill('solid', fgColor='BDD7EE')
            c.alignment = Alignment(horizontal='center')
            c.border    = BORDER
        return row + 1

    def data_row(row, label, *vals, bold=False):
        c = ws2.cell(row=row, column=1, value=label)
        c.font   = Font(name='Arial', bold=bold, size=10)
        c.border = BORDER
        fill_r   = FILL_WHITE if row % 2 == 0 else FILL_ALT
        c.fill   = fill_r
        for ci, v in enumerate(vals, 2):
            cell = ws2.cell(row=row, column=ci, value=v)
            cell.font      = Font(name='Arial', size=10)
            cell.alignment = Alignment(horizontal='center')
            cell.border    = BORDER
            cell.fill      = fill_r
            if isinstance(v, float):
                cell.number_format = '0.00'
        return row + 1

    # Gruplara ayır
    valid   = [r for r in rows if r.get('yas') is not None]
    erkek_r = [r for r in valid if str(r.get('cinsiyet','')).lower() == 'erkek']
    kadin_r = [r for r in valid if str(r.get('cinsiyet','')).lower() == 'kadin']

    def safe_mean(lst):
        return round(float(np.mean(lst)), 2) if lst else None

    def safe_std(lst):
        return round(float(np.std(lst)), 2) if lst else None

    def yas_mae(grp):
        errs = [abs(r['yas_tahmin'] - r['yas']) for r in grp
                if r.get('yas_tahmin') is not None and r.get('yas') is not None]
        return safe_mean(errs)

    def cin_acc(grp):
        if not grp:
            return None
        dogru = sum(1 for r in grp
                    if str(r.get('cinsiyet','')).lower() == str(r.get('tahmin_cin','')).lower())
        return round(dogru / len(grp) * 100, 1)

    r = 1
    r = sec_hdr(r, 'GENEL PERFORMANS OZETI')
    r = col_hdr(r, ['Metrik', 'Tum Hastalar', 'Erkek', 'Kadin'])

    r = data_row(r, 'Hasta Sayisi', len(valid), len(erkek_r), len(kadin_r), bold=True)
    r = data_row(r, 'Ort. Gercek Yas',
                 safe_mean([x['yas'] for x in valid]),
                 safe_mean([x['yas'] for x in erkek_r]),
                 safe_mean([x['yas'] for x in kadin_r]))
    r = data_row(r, 'Ort. Tahmin Yas',
                 safe_mean([x['yas_tahmin'] for x in valid if x.get('yas_tahmin')]),
                 safe_mean([x['yas_tahmin'] for x in erkek_r if x.get('yas_tahmin')]),
                 safe_mean([x['yas_tahmin'] for x in kadin_r if x.get('yas_tahmin')]))

    r += 1
    r = sec_hdr(r, 'YAS TAHMIN PERFORMANSI')
    r = col_hdr(r, ['Metrik', 'Tum Hastalar', 'Erkek', 'Kadin'])

    r = data_row(r, 'MAE (yil)',       yas_mae(valid),  yas_mae(erkek_r), yas_mae(kadin_r))
    all_errs = [abs(x['yas_tahmin'] - x['yas']) for x in valid
                if x.get('yas_tahmin') is not None and x.get('yas') is not None]
    r = data_row(r, 'Std Hata (yil)',  safe_std(all_errs), None, None)
    r = data_row(r, 'Max Hata (yil)',
                 round(max(all_errs), 2) if all_errs else None, None, None)
    r = data_row(r, 'Min Hata (yil)',
                 round(min(all_errs), 2) if all_errs else None, None, None)

    r += 1
    r = sec_hdr(r, 'CINSIYET TAHMIN PERFORMANSI')
    r = col_hdr(r, ['Metrik', 'Tum Hastalar', 'Erkek', 'Kadin'])

    r = data_row(r, 'Accuracy %',      cin_acc(valid),  cin_acc(erkek_r), cin_acc(kadin_r))
    cin_dogru_n = sum(1 for x in valid
                      if str(x.get('cinsiyet','')).lower() == str(x.get('tahmin_cin','')).lower())
    r = data_row(r, 'Dogru / Toplam',
                 f"{cin_dogru_n}/{len(valid)}",
                 f"{sum(1 for x in erkek_r if str(x.get('tahmin_cin','')).lower()=='erkek')}/{len(erkek_r)}",
                 f"{sum(1 for x in kadin_r if str(x.get('tahmin_cin','')).lower()=='kadin')}/{len(kadin_r)}")

    r += 1
    r = sec_hdr(r, 'HACIM TAHMIN PERFORMANSI (Uzman vs Pred)')
    r = col_hdr(r, ['Metrik', 'Tum Hastalar', 'Erkek', 'Kadin'])

    def vol_stats(grp, key_u, key_p):
        errs = [pct_error(x.get(key_p), x.get(key_u)) for x in grp
                if x.get(key_p) is not None and x.get(key_u) is not None]
        errs = [e for e in errs if e is not None]
        return safe_mean(errs)

    r = data_row(r, 'Ort. Uzman T (cm3)',
                 safe_mean([x['uzman_T'] for x in valid if x.get('uzman_T')]),
                 safe_mean([x['uzman_T'] for x in erkek_r if x.get('uzman_T')]),
                 safe_mean([x['uzman_T'] for x in kadin_r if x.get('uzman_T')]))
    r = data_row(r, 'Ort. Pred T (cm3)',
                 safe_mean([x['pred_T_cm3'] for x in valid if x.get('pred_T_cm3')]),
                 safe_mean([x['pred_T_cm3'] for x in erkek_r if x.get('pred_T_cm3')]),
                 safe_mean([x['pred_T_cm3'] for x in kadin_r if x.get('pred_T_cm3')]))
    r = data_row(r, 'Ort. R Hata %',
                 vol_stats(valid,   'uzman_R', 'pred_R_cm3'),
                 vol_stats(erkek_r, 'uzman_R', 'pred_R_cm3'),
                 vol_stats(kadin_r, 'uzman_R', 'pred_R_cm3'))
    r = data_row(r, 'Ort. L Hata %',
                 vol_stats(valid,   'uzman_L', 'pred_L_cm3'),
                 vol_stats(erkek_r, 'uzman_L', 'pred_L_cm3'),
                 vol_stats(kadin_r, 'uzman_L', 'pred_L_cm3'))
    r = data_row(r, 'Ort. T Hata %',
                 vol_stats(valid,   'uzman_T', 'pred_T_cm3'),
                 vol_stats(erkek_r, 'uzman_T', 'pred_T_cm3'),
                 vol_stats(kadin_r, 'uzman_T', 'pred_T_cm3'))

    r += 1
    r = sec_hdr(r, 'SEGMENTASYON METRIKLERI')
    r = col_hdr(r, ['Metrik', 'Ortalama', 'Std', None])

    for metric in ['dice', 'iou', 'precision', 'recall']:
        vals = [x[metric] for x in valid if x.get(metric) is not None]
        r = data_row(r, metric.capitalize(),
                     safe_mean(vals), safe_std(vals), None)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"[BILGI] Excel kaydedildi: {output_path}")


# ── Ana program ───────────────────────────────────────────

def main():
    os.makedirs(os.path.dirname(OUTPUT_EXCEL), exist_ok=True)

    print("=" * 65)
    print("MODELLER YUKLENIYOR")
    print("=" * 65)
    seg_model = tf.keras.models.load_model(
        SEG_MODEL_PATH, custom_objects={'f1_score': f1_score}
    )
    print("  + Segmentasyon modeli hazir")
    yas_model = joblib.load(MODEL_YAS_PATH)
    cin_model = joblib.load(MODEL_CIN_PATH)
    print(f"  + Yas modeli      ({yas_model.n_features_in_} ozellik: {FEATURE_COLS})")
    print(f"  + Cinsiyet modeli ({cin_model.n_features_in_} ozellik: {FEATURE_COLS})")

    assert yas_model.n_features_in_ == len(FEATURE_COLS), \
        "Yas modeli ozellik sayisi uyusmuyor! Asama2'yi yeniden calistirin."
    assert cin_model.n_features_in_ == len(FEATURE_COLS), \
        "Cinsiyet modeli ozellik sayisi uyusmuyor! Asama2'yi yeniden calistirin."

    excel_ref = load_excel_ref(EXCEL_REF_PATH)

    img_paths = sorted(glob(os.path.join(TEST_IMG_DIR, "*.nii*")))
    print(f"\n{len(img_paths)} test hastasi bulundu.\n")
    print("-" * 65)

    all_rows = []

    for idx, img_path in enumerate(img_paths, 1):
        fname     = os.path.basename(img_path)
        mask_path = os.path.join(TEST_MASK_DIR, fname)
        key       = parse_filename(fname)
        info      = excel_ref.get(key, {})

        print(f"[{idx:02d}/{len(img_paths)}] {fname:<22} "
              f"Yas:{str(info.get('yas','?')):>3}  "
              f"Cin:{str(info.get('cinsiyet','?')):<6}", end="  ->  ")

        seg = segment_sag_sol(
            seg_model, img_path,
            mask_path if os.path.exists(mask_path) else None
        )

        if seg is None:
            print("ATLANDI (sinus bulunamadi)")
            continue

        ml = predict_age_gender(seg, yas_model, cin_model)

        gy = info.get('yas')
        yas_hata_str = f"Hata:{abs(ml['yas_tahmin']-gy):.1f}yil" if gy else ""
        print(f"T={seg['pred_T_cm3']:.2f}cm3  "
              f"Yas={ml['yas_tahmin']:.1f} {yas_hata_str}  "
              f"Cin={ml['cinsiyet']}(%{ml['cin_guven_pct']:.0f})")

        all_rows.append({
            'hasta_adi'    : key,
            'yas'          : info.get('yas'),
            'yas_tahmin'   : ml['yas_tahmin'],
            'cinsiyet'     : info.get('cinsiyet', 'bilinmiyor'),
            'tahmin_cin'   : ml['cinsiyet'],
            'cin_guven_pct': ml['cin_guven_pct'],
            'uzman_R'      : info.get('uzman_R'),
            'uzman_L'      : info.get('uzman_L'),
            'uzman_T'      : info.get('uzman_T'),
            'pred_R_cm3'   : seg['pred_R_cm3'],
            'pred_L_cm3'   : seg['pred_L_cm3'],
            'pred_T_cm3'   : seg['pred_T_cm3'],
            'dice'         : seg.get('dice'),
            'iou'          : seg.get('iou'),
            'precision'    : seg.get('precision'),
            'recall'       : seg.get('recall'),
            'n_aktif_kesit': seg['n_aktif_kesit'],
            'slice_thick'  : seg['slice_thick'],
        })

    if not all_rows:
        print("\n[HATA] Hic sonuc uretilmedi.")
        return

    # Terminal özeti
    valid     = [r for r in all_rows if r['yas'] is not None]
    yas_hatalar = [abs(r['yas_tahmin'] - r['yas']) for r in valid]
    cin_dogru   = sum(1 for r in valid
                      if str(r['cinsiyet']).lower() == str(r['tahmin_cin']).lower())
    T_hatalar   = [pct_error(r['pred_T_cm3'], r['uzman_T'])
                   for r in valid if r.get('uzman_T')]
    T_hatalar   = [x for x in T_hatalar if x is not None]
    dice_vals   = [r['dice'] for r in all_rows if r.get('dice')]

    print(f"\n{'='*65}")
    print(f"SONUC OZETI  ({len(all_rows)} hasta islendi)")
    print(f"{'='*65}")
    if yas_hatalar:
        print(f"  Yas MAE          : {np.mean(yas_hatalar):.2f} yil  "
              f"+/- {np.std(yas_hatalar):.2f}")
    if valid:
        print(f"  Cinsiyet Acc     : %{cin_dogru/len(valid)*100:.1f}  "
              f"({cin_dogru}/{len(valid)})")
    if T_hatalar:
        print(f"  Ort. T Hata %    : %{np.mean(T_hatalar):.2f}")
    if dice_vals:
        print(f"  Ort. Dice        : {np.mean(dice_vals):.4f}")
    print(f"{'='*65}")

    write_excel(all_rows, OUTPUT_EXCEL)


if __name__ == "__main__":
    main()